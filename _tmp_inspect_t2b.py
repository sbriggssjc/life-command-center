"""Inspect master Excel charts for Round 31 Tier 2b items."""
import openpyxl
import sys
import warnings
warnings.filterwarnings("ignore")

GOV_PATH = r"C:\Users\scott\OneDrive - NorthMarq Capital, LLC\Team Briggs - Documents\Gv't Leased Research\Copy Government Master Document.xlsx"
DIA_PATH = r"C:\Users\scott\OneDrive - NorthMarq Capital, LLC\Team Briggs - Documents\Dialysis Research\Comps\Dialysis Comp Work MASTER.xlsx"


def chart_detail(ws, ch, idx):
    print(f"\n-- Chart {idx}  type={type(ch).__name__}")
    try:
        if ch.title and ch.title.tx and ch.title.tx.rich:
            for p in ch.title.tx.rich.paragraphs:
                for r in p.text:
                    if hasattr(r, 'value'):
                        print(f"   title: {r.value!r}")
                        break
    except Exception:
        pass
    try:
        ya = ch.y_axis
        if ya and ya.scaling:
            print(f"   y_axis: scaling={ya.scaling.min}..{ya.scaling.max}, fmt={ya.number_format}")
    except Exception:
        pass
    for j, s in enumerate(ch.series):
        nm = None
        try:
            if s.tx and s.tx.strRef and s.tx.strRef.f:
                ref = s.tx.strRef.f
                if "!" in ref:
                    sheet_part, cell_part = ref.split("!", 1)
                    sheet_clean = sheet_part.strip("'")
                    cell_clean = cell_part.replace("$", "")
                    target_ws = ws.parent[sheet_clean]
                    nm = target_ws[cell_clean].value
                else:
                    nm = ref
        except Exception:
            pass
        data_ref = None
        try:
            data_ref = s.val.numRef.f if s.val and s.val.numRef else None
        except Exception:
            pass
        print(f"   series[{j}]: label={nm!r}, data={data_ref}")


def dump(path, label, target):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    print(f"=== {label} ===")
    for sheet in target:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        print(f"\n## Sheet: {sheet}  (charts={len(ws._charts)})")
        for i, ch in enumerate(ws._charts):
            chart_detail(ws, ch, i)


if __name__ == "__main__":
    which = sys.argv[1] if len(sys.argv) > 1 else "both"
    if which in ("gov", "both"):
        dump(GOV_PATH, "GOV", ["All Charts", "SSA Charts", "Inventory"])
    if which in ("dia", "both"):
        dump(DIA_PATH, "DIA", ["Charts", "Market Size", "Available Coverage"])
