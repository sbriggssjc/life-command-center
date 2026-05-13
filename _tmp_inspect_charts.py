import openpyxl
import sys
from openpyxl.chart import BarChart, LineChart, ScatterChart, PieChart

GOV_PATH = r"C:\Users\scott\OneDrive - NorthMarq Capital, LLC\Team Briggs - Documents\Gv't Leased Research\Copy Government Master Document.xlsx"
DIA_PATH = r"C:\Users\scott\OneDrive - NorthMarq Capital, LLC\Team Briggs - Documents\Dialysis Research\Comps\Dialysis Comp Work MASTER.xlsx"


def chart_summary(ch, idx):
    out = {"index": idx, "type": type(ch).__name__}
    try:
        if ch.title is not None:
            for p in ch.title.tx.rich.paragraphs:
                for r in p.text:
                    if hasattr(r, 'value'):
                        out["title"] = r.value
                        break
                if "title" in out:
                    break
    except Exception:
        pass
    # Series titles
    series_names = []
    try:
        for s in ch.series:
            nm = None
            try:
                if s.tx and s.tx.strRef:
                    nm = s.tx.strRef.f
                elif s.tx and s.tx.v:
                    nm = s.tx.v
            except Exception:
                pass
            series_names.append(nm)
    except Exception:
        pass
    out["series"] = series_names
    # Axis labels
    try:
        if hasattr(ch, 'x_axis') and ch.x_axis and ch.x_axis.title:
            out["x_title"] = ch.x_axis.title.tx.rich.paragraphs[0].text[0].value
    except Exception:
        pass
    try:
        if hasattr(ch, 'y_axis') and ch.y_axis and ch.y_axis.title:
            out["y_title"] = ch.y_axis.title.tx.rich.paragraphs[0].text[0].value
    except Exception:
        pass
    return out


def dump(path, label, target_sheets=None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    print(f"=== {label} ===")
    for sheet in wb.sheetnames:
        if target_sheets and sheet not in target_sheets:
            continue
        ws = wb[sheet]
        if not ws._charts:
            continue
        print(f"\n## Sheet: {sheet}  (charts={len(ws._charts)})")
        for i, ch in enumerate(ws._charts):
            s = chart_summary(ch, i)
            print(f"  {s}")


if __name__ == "__main__":
    target_gov = ["All Charts", "SSA Charts"]
    target_dia = ["Charts", "Core Cap Chart", "Market Size"]
    if len(sys.argv) > 1 and sys.argv[1] == "gov":
        dump(GOV_PATH, "GOV", target_gov)
    elif len(sys.argv) > 1 and sys.argv[1] == "dia":
        dump(DIA_PATH, "DIA", target_dia)
    else:
        dump(GOV_PATH, "GOV", target_gov)
        dump(DIA_PATH, "DIA", target_dia)
